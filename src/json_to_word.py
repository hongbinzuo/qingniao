#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
将Discord JSON消息文件转换为Word文档
"""

import json
import sys
from datetime import datetime
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

def parse_timestamp(timestamp_str):
    """解析时间戳"""
    try:
        dt = datetime.fromisoformat(timestamp_str.replace('Z', '+00:00'))
        return dt.strftime('%Y-%m-%d %H:%M:%S')
    except:
        return timestamp_str

def format_message_content(content):
    """格式化消息内容，处理换行和特殊字符"""
    if not content:
        return ""
    # 替换换行符
    content = content.replace('\\n', '\n')
    return content

def json_to_word(json_file_path, output_file_path):
    """将JSON文件转换为Word文档"""
    print("正在读取JSON文件...")
    
    try:
        with open(json_file_path, 'r', encoding='utf-8') as f:
            data = json.load(f)
    except Exception as e:
        print(f"读取JSON文件失败: {e}")
        return False
    
    print("正在创建Word文档...")
    doc = Document()
    
    # 设置文档标题
    title = doc.add_heading('Discord 消息记录', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # 添加元信息
    guild = data.get('guild', {})
    channel = data.get('channel', {})
    
    info_para = doc.add_paragraph()
    info_para.add_run('服务器: ').bold = True
    info_para.add_run(guild.get('name', 'N/A'))
    info_para.add_run(' | 频道: ').bold = True
    info_para.add_run(channel.get('name', 'N/A'))
    
    exported_at = data.get('exportedAt', '')
    if exported_at:
        info_para = doc.add_paragraph()
        info_para.add_run('导出时间: ').bold = True
        info_para.add_run(exported_at)
    
    doc.add_paragraph()  # 空行
    
    # 处理消息
    messages = data.get('messages', [])
    total_messages = len(messages)
    print(f"找到 {total_messages} 条消息，正在处理...")
    
    current_date = None
    
    for i, msg in enumerate(messages):
        if (i + 1) % 100 == 0:
            print(f"已处理 {i + 1}/{total_messages} 条消息...")
        
        # 获取消息信息
        timestamp = msg.get('timestamp', '')
        author = msg.get('author', {})
        content = msg.get('content', '')
        msg_type = msg.get('type', 'Default')
        
        # 解析时间戳
        try:
            msg_time = parse_timestamp(timestamp)
            msg_date = msg_time.split(' ')[0] if msg_time else None
        except:
            msg_date = None
            msg_time = timestamp
        
        # 如果是新的一天，添加日期分隔
        if msg_date and msg_date != current_date:
            current_date = msg_date
            date_heading = doc.add_heading(msg_date, level=2)
            date_heading.style.font.size = Pt(14)
        
        # 添加消息段落
        msg_para = doc.add_paragraph()
        
        # 作者名称（加粗）
        author_name = author.get('nickname') or author.get('name', 'Unknown')
        author_run = msg_para.add_run(f"{author_name}")
        author_run.bold = True
        author_run.font.size = Pt(11)
        
        # 时间戳
        time_run = msg_para.add_run(f" [{msg_time}]")
        time_run.font.size = Pt(9)
        time_run.font.color.rgb = RGBColor(128, 128, 128)
        
        # 消息类型（如果不是默认类型）
        if msg_type != 'Default':
            type_run = msg_para.add_run(f" [{msg_type}]")
            type_run.font.size = Pt(9)
            type_run.font.color.rgb = RGBColor(100, 100, 200)
        
        # 消息内容
        if content:
            content_para = doc.add_paragraph(format_message_content(content))
            content_para.style.font.size = Pt(10)
        
        # 处理附件
        attachments = msg.get('attachments', [])
        if attachments:
            for att in attachments:
                att_para = doc.add_paragraph()
                att_run = att_para.add_run(f"📎 附件: {att.get('fileName', 'Unknown')}")
                att_run.font.size = Pt(9)
                att_run.font.color.rgb = RGBColor(0, 100, 200)
        
        # 处理回复/引用（如果有）
        if msg.get('referencedMessage'):
            ref_para = doc.add_paragraph()
            ref_run = ref_para.add_run("↩ 回复消息")
            ref_run.font.size = Pt(9)
            ref_run.font.color.rgb = RGBColor(150, 150, 150)
            ref_run.italic = True
        
        # 添加小间距
        doc.add_paragraph()
    
    print("正在保存Word文档...")
    try:
        doc.save(output_file_path)
        print(f"成功！Word文档已保存到: {output_file_path}")
        return True
    except Exception as e:
        print(f"保存Word文档失败: {e}")
        return False

def main():
    input_file = r"C:\Users\zuoho\Downloads\discord-msg\Direct Messages - 知更 [1361229030715687024].json"
    output_file = r"C:\Users\zuoho\Downloads\discord-msg\Direct Messages - 知更 [1361229030715687024].docx"
    
    print("=" * 80)
    print("Discord JSON 转 Word 工具")
    print("=" * 80)
    print()
    print(f"输入文件: {input_file}")
    print(f"输出文件: {output_file}")
    print()
    
    success = json_to_word(input_file, output_file)
    
    if success:
        print()
        print("=" * 80)
        print("转换完成！")
        print("=" * 80)
    else:
        print()
        print("=" * 80)
        print("转换失败，请检查错误信息")
        print("=" * 80)
        sys.exit(1)

if __name__ == "__main__":
    main()




